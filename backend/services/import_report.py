"""Формирование XLSX-отчёта о проблемах импорта каталога.

Отчёт читает человек: сначала шапка с итогами, затем плоский список
проблем с указанием строки, поля и причины. Ошибка блокирует строку,
предупреждение — нет.
"""

from __future__ import annotations

from io import BytesIO

from app.services.excel_import import (
    SEVERITY_ERROR,
    SEVERITY_WARNING,
    CatalogImportResult,
)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = ("Строка", "Поле", "Проблема", "Значение", "Уровень")
COLUMN_WIDTHS = (10, 24, 44, 40, 14)

_SEVERITY_LABELS = {
    SEVERITY_ERROR: "Ошибка",
    SEVERITY_WARNING: "Предупреждение",
}

_ERROR_FILL = PatternFill("solid", fgColor="FDE7E9")
_WARNING_FILL = PatternFill("solid", fgColor="FFF4E5")
_HEADER_FILL = PatternFill("solid", fgColor="E8EEF7")


def generate_import_report(result: CatalogImportResult) -> BytesIO:
    """Собирает XLSX с проблемами валидации: одна строка — одна проблема."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Проблемы импорта"

    summary = result.to_dict()["summary"]
    sheet.append([
        f"Проверено строк: {summary['total_rows']}",
        f"Готовы к импорту: {summary['valid_rows']}",
        f"Ошибок: {summary['error_rows']}",
        f"Предупреждений: {summary['warning_rows']}",
    ])
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    sheet.append([])
    sheet.append(list(HEADERS))
    header_row = sheet.max_row
    for cell in sheet[header_row]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL

    for issue in result.issues:
        sheet.append([
            issue.row if issue.row is not None else "—",
            issue.field or "—",
            issue.problem,
            issue.value if issue.value is not None else "—",
            _SEVERITY_LABELS.get(issue.severity, issue.severity),
        ])
        fill = _ERROR_FILL if issue.severity == SEVERITY_ERROR else _WARNING_FILL
        for cell in sheet[sheet.max_row]:
            cell.fill = fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for index, width in enumerate(COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
