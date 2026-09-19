"""Импорт каталогов из XLSX файлов."""

from typing import Any
from io import BytesIO

try:
    from openpyxl import load_workbook
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


class CatalogImportResult:
    """Результат импорта каталога."""
    
    def __init__(self, success: bool, data: list[dict], errors: list[str] = None):
        self.success = success
        self.data = data
        self.errors = errors or []
    
    def to_dict(self) -> dict:
        return {
            "success": self.success,
            "data": self.data,
            "errors": self.errors,
            "count": len(self.data)
        }


def parse_catalog_file(file: BytesIO, catalog_type: str) -> CatalogImportResult:
    """Парсинг файла каталога (вузов или продуктов).
    
    Args:
        file: BytesIO с содержимым файла
        catalog_type: 'universities' или 'products'
    
    Returns:
        CatalogImportResult с данными и ошибками
    """
    if not OPENPYXL_AVAILABLE:
        return CatalogImportResult(
            success=False,
            data=[],
            errors=["openpyxl не установлен. Установите: pip install openpyxl"]
        )
    
    try:
        wb = load_workbook(file)
        ws = wb.active
        
        if catalog_type == "universities":
            return _parse_universities(ws)
        elif catalog_type == "products":
            return _parse_products(ws)
        else:
            return CatalogImportResult(
                success=False,
                data=[],
                errors=[f"Неизвестный тип каталога: {catalog_type}"]
            )
    except Exception as e:
        return CatalogImportResult(
            success=False,
            data=[],
            errors=[f"Ошибка чтения файла: {str(e)}"]
        )


def _parse_universities(worksheet) -> CatalogImportResult:
    """Парсинг листа с вузами.
    
    Ожидаемые колонки:
    - Название (обязательно)
    - Город
    - Контактное лицо
    - Email
    - Телефон
    """
    data = []
    errors = []
    
    # Проверяем заголовки
    headers = []
    for cell in worksheet[1]:
        headers.append(cell.value)
    
    if not headers or all(h is None for h in headers):
        return CatalogImportResult(
            success=False,
            data=[],
            errors=["Файл пуст или отсутствуют заголовки"]
        )
    
    # Нормализуем заголовки
    header_map = {
        "название": "name",
        "город": "city",
        "контактное лицо": "contact_person",
        "email": "contact_email",
        "телефон": "contact_phone",
        "name": "name",
        "city": "city",
        "contact person": "contact_person",
        "contact email": "contact_email",
        "contact phone": "contact_phone",
    }
    
    column_indices = {}
    for idx, header in enumerate(headers):
        if header and str(header).lower().strip() in header_map:
            column_indices[header_map[str(header).lower().strip()]] = idx
    
    if "name" not in column_indices:
        return CatalogImportResult(
            success=False,
            data=[],
            errors=["Не найдена обязательная колонка 'Название'"]
        )
    
    # Парсим данные
    for row_idx, row in enumerate(worksheet.iter_rows(min_row=2), start=2):
        if not any(cell.value for cell in row):
            continue  # Пропускаем пустые строки
        
        try:
            university = {
                "name": row[column_indices["name"]].value,
                "city": row.get(column_indices.get("city", 1), None).value if "city" in column_indices else None,
                "contact_person": row.get(column_indices.get("contact_person", 2), None).value if "contact_person" in column_indices else None,
                "contact_email": row.get(column_indices.get("contact_email", 3), None).value if "contact_email" in column_indices else None,
                "contact_phone": row.get(column_indices.get("contact_phone", 4), None).value if "contact_phone" in column_indices else None,
            }
            
            # Валидация
            if not university["name"]:
                errors.append(f"Строка {row_idx}: отсутствует название")
                continue
            
            data.append(university)
        except Exception as e:
            errors.append(f"Строка {row_idx}: {str(e)}")
    
    return CatalogImportResult(success=True, data=data, errors=errors)


def _parse_products(worksheet) -> CatalogImportResult:
    """Парсинг листа с продуктами.
    
    Ожидаемые колонки:
    - Название (обязательно)
    - Направление
    """
    data = []
    errors = []
    
    # Проверяем заголовки
    headers = []
    for cell in worksheet[1]:
        headers.append(cell.value)
    
    if not headers or all(h is None for h in headers):
        return CatalogImportResult(
            success=False,
            data=[],
            errors=["Файл пуст или отсутствуют заголовки"]
        )
    
    # Нормализуем заголовки
    header_map = {
        "название": "name",
        "направление": "direction",
        "name": "name",
        "direction": "direction",
    }
    
    column_indices = {}
    for idx, header in enumerate(headers):
        if header and str(header).lower().strip() in header_map:
            column_indices[header_map[str(header).lower().strip()]] = idx
    
    if "name" not in column_indices:
        return CatalogImportResult(
            success=False,
            data=[],
            errors=["Не найдена обязательная колонка 'Название'"]
        )
    
    # Парсим данные
    for row_idx, row in enumerate(worksheet.iter_rows(min_row=2), start=2):
        if not any(cell.value for cell in row):
            continue  # Пропускаем пустые строки
        
        try:
            product = {
                "name": row[column_indices["name"]].value,
                "direction": row.get(column_indices.get("direction", 1), None).value if "direction" in column_indices else None,
            }
            
            # Валидация
            if not product["name"]:
                errors.append(f"Строка {row_idx}: отсутствует название")
                continue
            
            data.append(product)
        except Exception as e:
            errors.append(f"Строка {row_idx}: {str(e)}")
    
    return CatalogImportResult(success=True, data=data, errors=errors)
