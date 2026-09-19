"""Экспорт отчётов в XLSX, XLS и PDF форматы."""

from io import BytesIO
from datetime import datetime

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

try:
    from xlwt import Workbook as XlsWorkbook, easyxf
    XLWT_AVAILABLE = True
except ImportError:
    XLWT_AVAILABLE = False

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False


def generate_xlsx(interactions: list[dict]) -> BytesIO:
    """Генерация XLSX отчёта по взаимодействиям."""
    if not OPENPYXL_AVAILABLE:
        raise ImportError("openpyxl не установлен. Установите: pip install openpyxl")
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Отчёт по взаимодействиям"
    
    # Заголовки
    headers = ["ID", "Вуз", "Продукт", "Этап", "Номер договора", "Дата договора", 
               "Специалист РТК", "КАМ", "Специалист вуза", "Заметки", "Активен"]
    
    # Стиль заголовков
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Запись заголовков
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    # Запись данных
    for row_num, interaction in enumerate(interactions, 2):
        ws.cell(row=row_num, column=1, value=interaction.get("id"))
        ws.cell(row=row_num, column=2, value=interaction.get("university_name"))
        ws.cell(row=row_num, column=3, value=interaction.get("product_name"))
        ws.cell(row=row_num, column=4, value=interaction.get("stage_name"))
        ws.cell(row=row_num, column=5, value=interaction.get("contract_number"))
        ws.cell(row=row_num, column=6, value=interaction.get("contract_date"))
        ws.cell(row=row_num, column=7, value=interaction.get("rkn_specialist_name"))
        ws.cell(row=row_num, column=8, value=interaction.get("assigned_kam_name"))
        ws.cell(row=row_num, column=9, value=interaction.get("university_specialist"))
        ws.cell(row=row_num, column=10, value=interaction.get("notes"))
        ws.cell(row=row_num, column=11, value="Да" if interaction.get("is_active") else "Нет")
        
        # Применение границ к ячейкам данных
        for col_num in range(1, 12):
            cell = ws.cell(row=row_num, column=col_num)
            cell.border = thin_border
    
    # Автоподбор ширины колонок
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column].width = adjusted_width
    
    # Сохранение в BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_xls(interactions: list[dict]) -> BytesIO:
    """Генерация XLS отчёта по взаимодействиям."""
    if not XLWT_AVAILABLE:
        raise ImportError("xlwt не установлен. Установите: pip install xlwt")
    
    wb = XlsWorkbook()
    ws = wb.add_sheet("Отчёт по взаимодействиям")
    
    # Заголовки
    headers = ["ID", "Вуз", "Продукт", "Этап", "Номер договора", "Дата договора", 
               "Специалист РТК", "КАМ", "Специалист вуза", "Заметки", "Активен"]
    
    # Стиль заголовков
    header_style = easyxf(
        'font: bold on, color white; align: horiz center, vert centre; '
        'pattern: pattern solid, fore_color blue_grey'
    )
    
    # Запись заголовков
    for col_num, header in enumerate(headers):
        ws.write(0, col_num, header, header_style)
    
    # Запись данных
    for row_num, interaction in enumerate(interactions, 1):
        ws.write(row_num, 0, interaction.get("id"))
        ws.write(row_num, 1, interaction.get("university_name"))
        ws.write(row_num, 2, interaction.get("product_name"))
        ws.write(row_num, 3, interaction.get("stage_name"))
        ws.write(row_num, 4, interaction.get("contract_number"))
        ws.write(row_num, 5, interaction.get("contract_date"))
        ws.write(row_num, 6, interaction.get("rkn_specialist_name"))
        ws.write(row_num, 7, interaction.get("assigned_kam_name"))
        ws.write(row_num, 8, interaction.get("university_specialist"))
        ws.write(row_num, 9, interaction.get("notes"))
        ws.write(row_num, 10, "Да" if interaction.get("is_active") else "Нет")
    
    # Сохранение в BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_pdf(interactions: list[dict]) -> BytesIO:
    """Генерация PDF отчёта по взаимодействиям."""
    if not REPORTLAB_AVAILABLE:
        raise ImportError("reportlab не установлен. Установите: pip install reportlab")
    
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=letter)
    
    # Стиль для заголовков
    styles = getSampleStyleSheet()
    title_style = styles["Title"]
    title_style.alignment = 1  # center
    
    # Данные для таблицы
    headers = ["ID", "Вуз", "Продукт", "Этап", "Договор", "Специалист РТК", "КАМ", "Активен"]
    data = [headers]
    
    for interaction in interactions:
        row = [
            str(interaction.get("id", "")),
            interaction.get("university_name", ""),
            interaction.get("product_name", ""),
            interaction.get("stage_name", ""),
            f"{interaction.get('contract_number', '')} ({interaction.get('contract_date', '')})",
            interaction.get("rkn_specialist_name", ""),
            interaction.get("assigned_kam_name", ""),
            "Да" if interaction.get("is_active") else "Нет"
        ]
        data.append(row)
    
    # Создание таблицы
    table = Table(data, colWidths=[0.5*inch, 1.5*inch, 1.2*inch, 1.2*inch, 1.2*inch, 1.2*inch, 1*inch, 0.8*inch])
    
    # Стиль таблицы
    table_style = TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
    ])
    table.setStyle(table_style)
    
    # Заголовок документа
    title = Paragraph(f"Отчёт по взаимодействиям ({datetime.now().strftime('%d.%m.%Y %H:%M')})", title_style)
    
    # Построение документа
    elements = [title, table]
    doc.build(elements)
    
    output.seek(0)
    return output
