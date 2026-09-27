"""Обезличивает выгрузку пользователей кейсодержателя для публикации.

Читает «Загрузка пользователей.xlsx», заменяет поля с персональными данными
(СНИЛС, паспорт, адрес регистрации) на сгенерированные заглушки и сохраняет
результат как `docs/examples/users_b2c.xlsx`. ФИО, телефон и email остаются
тестовыми, реальных ПДн в файле нет.

Запуск из корня репозитория:

    python scripts/anonymize_case_data.py <путь к «Загрузка пользователей.xlsx»>
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parent.parent
TARGET = ROOT / "docs" / "examples" / "users_b2c.xlsx"

# Колонка → заглушка. Адрес разнесён по колонкам выгрузки, поэтому вместо
# одной строки подставляется соответствующая часть, а индекс — отдельно.
REPLACEMENTS: dict[str, str] = {
    "СНИЛС": "000-000-000 00",
    "Серия паспорта": "0000",
    "Номер паспорта": "000000",
    "Кем выдан паспорт": "ОТДЕЛОМ УФМС РОССИИ",
    "Дата выдачи паспорта": "01.01.2020",
    "Код подразделения": "000-000",
    "Регион регистрации": "г. Москва",
    "Населенный пункт регистрации": "г. Москва",
    "Улица регистрации": "ул. Примерная",
    "Дом регистрации": "д. 1",
    "Квартира регистрации": "кв. 1",
    "Индекс регистрации": "000000",
}


def anonymize(sheet: Worksheet) -> list[str]:
    headers = {cell.value: cell.column for cell in sheet[1] if cell.value}
    touched: list[str] = []
    for column, placeholder in REPLACEMENTS.items():
        index = headers.get(column)
        if index is None:
            continue
        touched.append(column)
        for row in range(2, sheet.max_row + 1):
            sheet.cell(row=row, column=index).value = placeholder
    return touched


def main() -> int:
    if len(sys.argv) != 2:
        print(__doc__)
        return 2
    source = Path(sys.argv[1])
    if not source.exists():
        print(f"Нет файла {source}")
        return 1

    workbook = openpyxl.load_workbook(source)
    touched = anonymize(workbook.worksheets[0])
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(TARGET)
    print(f"Записано {TARGET.relative_to(ROOT)}; обезличено колонок: {len(touched)}")
    print("Поля:", ", ".join(touched))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
