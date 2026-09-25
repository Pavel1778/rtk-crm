"""Тесты отчёта о проблемах импорта каталога.

Отчёт — часть импорта, а не отдельная функция: те же строки, что в
предпросмотре, но с указанием строки, поля и уровня проблемы. Эти тесты
фиксируют разницу между ошибкой (строка не импортируется) и
предупреждением (импортируется, но менеджер должен знать).
"""

from __future__ import annotations

from io import BytesIO
from uuid import uuid4

import openpyxl
import pytest
from app.db.session import SessionLocal, create_tables
from app.main import app
from app.models.entities import User
from app.models.enums import UserRole
from app.services.excel_import import (
    SEVERITY_ERROR,
    SEVERITY_WARNING,
    parse_catalog_file,
)
from app.services.import_report import HEADERS, generate_import_report
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook
from sqlalchemy import select

UNIVERSITY_HEADERS = ["Название", "Город", "Контактное лицо", "Email", "Телефон"]

ADMIN_EMAIL = "admin@rtk.ru"
ADMIN_PASSWORD = "admin123"


@pytest.fixture(autouse=True)
async def _prepare_db():
    """Таблицы + администратор: тесты не должны зависеть от внешнего сида."""
    await create_tables()
    from app.auth.security import hash_password
    from app.models.entities import User
    from app.models.enums import UserRole

    async with SessionLocal() as session:
        existing = await session.scalar(
            select(User).where(User.email == ADMIN_EMAIL)
        )
        if existing is None:
            session.add(
                User(
                    email=ADMIN_EMAIL,
                    full_name="Администратор",
                    role=UserRole.ADMIN,
                    is_admin=True,
                    hashed_password=hash_password(ADMIN_PASSWORD),
                )
            )
            await session.commit()


def _xlsx(rows: list[list[object]]) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(UNIVERSITY_HEADERS)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _parse(rows: list[list[object]]):
    return parse_catalog_file(_xlsx(rows), "universities", "catalog.xlsx")


def test_report_summary_counts_valid_error_and_warning_rows() -> None:
    result = _parse([
        ["МГУ", "Москва", "Иван", "ivan@example.com", "+7 900 000-00-00"],
        ["", "Москва", "", "", ""],
        ["РТУ МИРЭА", "Москва", "", "not-an-email", ""],
        ["МГУ", "Москва", "", "", ""],
    ])

    summary = result.to_dict()["summary"]
    assert summary["total_rows"] == 4
    assert summary["valid_rows"] == 3
    assert summary["error_rows"] == 1
    assert summary["warning_rows"] == 2


def test_empty_required_field_is_error_row() -> None:
    result = _parse([["", "Москва", "", "", ""]])

    assert result.data == []
    assert result.errors == ["Строка 2: отсутствует название"]
    issue = result.issues[0]
    assert issue.severity == SEVERITY_ERROR
    assert issue.row == 2
    assert issue.field == "Название"


def test_duplicate_name_is_warning_and_row_still_imported() -> None:
    result = _parse([
        ["МГУ", "Москва", "", "", ""],
        ["МГУ", "Москва", "", "", ""],
    ])

    assert len(result.data) == 2
    warnings = [issue for issue in result.issues if issue.severity == SEVERITY_WARNING]
    assert len(warnings) == 1
    assert warnings[0].row == 3
    assert warnings[0].problem == "Дубль в файле (строка 2)"


def test_invalid_email_is_warning() -> None:
    result = _parse([["СПбГУ", "СПб", "", "student(at)example.com", ""]])

    assert len(result.data) == 1
    issue = result.issues[0]
    assert issue.severity == SEVERITY_WARNING
    assert issue.field == "Email"
    assert issue.problem == "Некорректный email"
    assert issue.value == "student(at)example.com"


def test_missing_required_column_reports_single_issue() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Город"])
    sheet.append(["Москва"])
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    result = parse_catalog_file(output, "universities", "catalog.xlsx")

    assert result.success is False
    assert result.issues[0].field == "Название"
    assert result.issues[0].row is None


def test_blank_rows_are_ignored_entirely() -> None:
    result = _parse([
        ["МГУ", "Москва", "", "", ""],
        [None, None, None, None, None],
    ])

    summary = result.to_dict()["summary"]
    assert summary["total_rows"] == 1
    assert summary["error_rows"] == 0
    assert summary["warning_rows"] == 0


def test_xlsx_report_has_headers_and_one_row_per_issue() -> None:
    result = _parse([
        ["", "Москва", "", "", ""],
        ["СПбГУ", "СПб", "", "bad-email", ""],
    ])

    stream = generate_import_report(result)
    sheet = openpyxl.load_workbook(stream).active

    header_row = next(
        index
        for index, row in enumerate(sheet.iter_rows(values_only=True), start=1)
        if row and tuple(row)[: len(HEADERS)] == HEADERS
    )
    data = list(sheet.iter_rows(min_row=header_row + 1, values_only=True))
    assert len(data) == len(result.issues) == 2
    assert data[0][4] == "Ошибка"
    assert data[1][4] == "Предупреждение"


def test_report_marks_fully_valid_file_without_issues() -> None:
    result = _parse([["МГУ", "Москва", "Иван", "ivan@example.com", "+7"]])

    assert result.issues == []
    summary = result.to_dict()["summary"]
    assert summary["error_count"] == 0
    assert summary["warning_count"] == 0


# --- Контракт эндпоинта: интерфейс зависит от этих деталей ---


def _report_client() -> AsyncClient:
    return AsyncClient(transport=ASGITransport(app=app), base_url="http://test")


async def _auth_headers(client: AsyncClient) -> dict[str, str]:
    login = await client.post(
        "/api/auth/login",
        json={"email": "admin@rtk.ru", "password": "admin123"},
    )
    assert login.status_code == 200
    return {"Authorization": f"Bearer {login.json()['access_token']}"}


async def test_report_endpoint_requires_authentication() -> None:
    async with _report_client() as client:
        response = await client.post(
            "/api/catalogs/import/report",
            params={"catalog_type": "universities"},
            files={"file": ("catalog.xlsx", _xlsx([["МГУ", "Москва", "", "", ""]]).getvalue())},
        )

    assert response.status_code == 401


async def test_report_endpoint_returns_xlsx_with_counts_in_headers() -> None:
    async with _report_client() as client:
        headers = await _auth_headers(client)
        response = await client.post(
            "/api/catalogs/import/report",
            params={"catalog_type": "universities"},
            files={
                "file": (
                    "catalog.xlsx",
                    _xlsx([
                        ["МГУ", "Москва", "", "", ""],
                        ["", "Москва", "", "", ""],
                    ]).getvalue(),
                )
            },
            data={"mapping": "{}"},
            headers=headers,
        )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "import-report-" in response.headers["content-disposition"]
    assert response.headers["x-import-error-rows"] == "1"
    assert response.headers["x-import-warning-rows"] == "0"

    sheet = openpyxl.load_workbook(BytesIO(response.content)).active
    values = [cell for row in sheet.iter_rows(values_only=True) for cell in row if cell]
    assert "Ошибок: 1" in values


async def test_execute_imports_valid_rows_and_reports_skipped() -> None:
    async with _report_client() as client:
        headers = await _auth_headers(client)
        name = f"Только валидные {uuid4().hex[:8]}"
        response = await client.post(
            "/api/catalogs/import/execute",
            params={"catalog_type": "universities"},
            files={
                "file": (
                    "catalog.xlsx",
                    _xlsx([
                        [name, "Москва", "Иван", "ivan@example.com", "+7"],
                        ["", "Москва", "", "", ""],
                    ]).getvalue(),
                )
            },
            data={
                "mapping": "{}",
            },
            headers=headers,
        )

    assert response.status_code == 200
    body = response.json()
    # Одна строка импортируется, строка с пустым названием отсекается валидацией.
    assert body["created"] == 1
    assert body["total"] == 1
    assert body["errors"] == ["Строка 3: отсутствует название"]


async def test_report_endpoint_rejects_non_excel_and_unknown_catalog() -> None:
    async with _report_client() as client:
        headers = await _auth_headers(client)
        wrong_type = await client.post(
            "/api/catalogs/import/report",
            params={"catalog_type": "unknown"},
            files={"file": ("catalog.xlsx", _xlsx([["МГУ", "Москва", "", "", ""]]).getvalue())},
            data={"mapping": "{}"},
            headers=headers,
        )
        wrong_format = await client.post(
            "/api/catalogs/import/report",
            params={"catalog_type": "universities"},
            files={"file": ("catalog.json", b"[]")},
            data={"mapping": "{}"},
            headers=headers,
        )

    assert wrong_type.status_code == 400
    assert wrong_format.status_code == 400


def _users_xlsx(rows: list[list[object]]) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Фамилия", "Имя", "Отчество", "Email", "СНИЛС"])
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


async def test_execute_imports_users_from_case_holder_export() -> None:
    """Выгрузка кейсодержателя создаёт учётные записи без ПДн."""
    async with _report_client() as client:
        headers = await _auth_headers(client)
        email = f"case-holder-{uuid4().hex[:8]}@example.com"
        response = await client.post(
            "/api/catalogs/import/execute",
            params={"catalog_type": "users"},
            files={
                "file": (
                    "Загрузка пользователей.xlsx",
                    _users_xlsx([
                        ["Черепанова", "Светлана", "Васильевна", email, "123-456-789 00"],
                    ]).getvalue(),
                )
            },
            data={"mapping": "{}"},
            headers=headers,
        )

    assert response.status_code == 200
    body = response.json()
    assert body["created"] == 1
    assert body["total"] == 1

    async with SessionLocal() as session:
        user = await session.scalar(select(User).where(User.email == email))
        assert user is not None
        assert user.full_name == "Черепанова Светлана Васильевна"
        assert user.role == UserRole.USER
        # Пароль из выгрузки не берётся: хеш временного пароля не пустой.
        assert user.hashed_password


async def test_execute_rejects_unknown_catalog_type() -> None:
    async with _report_client() as client:
        headers = await _auth_headers(client)
        response = await client.post(
            "/api/catalogs/import/execute",
            params={"catalog_type": "unknown"},
            files={"file": ("catalog.xlsx", _xlsx([["МГУ", "Москва", "", "", ""]]).getvalue())},
            data={"mapping": "{}"},
            headers=headers,
        )

    assert response.status_code == 400
