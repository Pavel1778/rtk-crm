from io import BytesIO

import xlwt
from openpyxl import Workbook

from backend.services.excel_import import parse_catalog_file, parse_catalog_json

HEADERS = ["Название", "Город", "Контактное лицо", "Email", "Телефон"]
ROW = ["МГУ им. Ленина №5", "Москва", "Иван Петров", "ivan@example.com", "+7 900 000-00-00"]


def _xlsx_bytes() -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADERS)
    sheet.append(ROW)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _xls_bytes() -> BytesIO:
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Вузы")
    for column, value in enumerate(HEADERS):
        sheet.write(0, column, value)
    for column, value in enumerate(ROW):
        sheet.write(1, column, value)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def test_xlsx_and_xls_have_identical_normalized_data() -> None:
    xlsx = parse_catalog_file(_xlsx_bytes(), "universities", "sample.xlsx")
    xls = parse_catalog_file(_xls_bytes(), "universities", "sample.xls")

    assert xlsx.success is True
    assert xls.success is True
    assert xlsx.headers == xls.headers == HEADERS
    assert xlsx.data == xls.data == [
        {
            "name": ROW[0],
            "city": ROW[1],
            "contact_person": ROW[2],
            "contact_email": ROW[3],
            "contact_phone": ROW[4],
        }
    ]


def test_optional_columns_do_not_break_xlsx_import() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Название"])
    sheet.append(["СПбГУ"])
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    result = parse_catalog_file(output, "universities", "sample.xlsx")

    assert result.success is True
    assert result.data[0]["name"] == "СПбГУ"
    assert result.data[0]["city"] is None


def test_json_import_supports_explicit_column_mapping() -> None:
    content = (
        '[{"Название ВУЗа": "КФУ", "Населенный пункт": "Казань"}]'
    ).encode()

    result = parse_catalog_json(
        content,
        "universities",
        {"name": "Название ВУЗа", "city": "Населенный пункт"},
    )

    assert result.success is True
    assert result.data == [
        {
            "name": "КФУ",
            "city": "Казань",
            "contact_person": None,
            "contact_email": None,
            "contact_phone": None,
        }
    ]


def test_json_import_rejects_malformed_payload_with_path() -> None:
    result = parse_catalog_json(b'{"data": "not-an-array"}', "universities")

    assert result.success is False
    assert "data" in result.errors[0]
    assert "Ошибка структуры JSON" in result.errors[0]


def test_json_import_rejects_empty_record_object() -> None:
    result = parse_catalog_json(b"[{}]", "universities")

    assert result.success is False
    assert "пустой объект" in result.errors[0]


def test_json_import_accepts_wrapped_data_object() -> None:
    content = b'{"data": [{"name": "\\u041c\\u0413\\u0423"}]}'

    result = parse_catalog_json(content, "universities")

    assert result.success is True
    assert result.data[0]["name"] == "МГУ"


def test_json_import_reports_required_name_missing() -> None:
    payload = b'[{"city": "\\u041a\\u0430\\u0437\\u0430\\u043d\\u044c"}]'
    result = parse_catalog_json(payload, "universities")

    assert result.success is False
    assert any("Название" in error for error in result.errors)


def _vendors_xlsx() -> BytesIO:
    """Мини-аналог «Вендоры.xlsx» кейсодержателя: продукт вместо «Названия»."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Компания", "Продукт", "ФИО", "Телефон", "Почта", "Способ связи"])
    sheet.append(
        ["ООО «Базис»", "«Базис Dynamix»", "Иванов И. И.", "+7 900 111-22-33",
         "ivanov@example.ru", "Чат в ТГ"]
    )
    sheet.append(
        ["ООО «РТК ИТ»", "«Нейрошлюз»", "Козлов М. И.", "+7 900 777-88-99",
         "kozlov@example.ru", "Почта"]
    )
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def test_products_import_accepts_vendor_catalog() -> None:
    """Колонка «Продукт» из «Вендоры.xlsx» распознаётся как название продукта."""
    result = parse_catalog_file(_vendors_xlsx(), "products", "Вендоры.xlsx")

    assert result.success is True
    assert [item["name"] for item in result.data] == [
        "«Базис Dynamix»",
        "«Нейрошлюз»",
    ]
    # Контакты вендоров в каталог продуктов не переносятся.
    assert all(set(item) == {"name", "direction"} for item in result.data)


def test_products_import_accepts_course_column_from_json() -> None:
    """Колонка «Курс» из «Данные оплат.json» тоже описывает ИТ-продукт."""
    payload = (
        b'[{"\\u041a\\u0443\\u0440\\u0441": "\\u041f\\u0440\\u043e\\u043c\\u043f\\u0442"}]'
    )

    result = parse_catalog_json(payload, "products")

    assert result.success is True
    assert result.data[0]["name"] == "Промпт"


def test_json_import_skips_leading_null_stub() -> None:
    """Выгрузки кейсодержателя начинаются с null: строка не должна ломать файл."""
    payload = b'[null, {"name": "\\u041c\\u0413\\u0423"}]'

    result = parse_catalog_json(payload, "universities")

    assert result.success is True
    assert result.data == [{"name": "МГУ", "city": None, "contact_person": None,
                            "contact_email": None, "contact_phone": None}]


def _users_xlsx() -> BytesIO:
    """Мини-аналог «Загрузка пользователей.xlsx» с ПДн-колонками."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        ["Фамилия", "Имя", "Отчество", "Email", "СНИЛС",
         "Серия паспорта", "Номер паспорта", "Улица регистрации"]
    )
    sheet.append(
        ["Черепанова", "Светлана", "Васильевна", "cherepanona.s@test.ru",
         "123-456-789 00", "4509", "123456", "ул. Ленина"]
    )
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def test_users_import_collects_only_account_fields() -> None:
    """Импорт пользователей не переносит ПДн из выгрузки (152-ФЗ)."""
    result = parse_catalog_file(_users_xlsx(), "users", "Загрузка пользователей.xlsx")

    assert result.success is True
    assert result.data == [
        {
            "email": "cherepanona.s@test.ru",
            "full_name": "Черепанова Светлана Васильевна",
            "role": None,
        }
    ]
    serialised = str(result.data)
    for sensitive in ("СНИЛС", "123-456-789", "4509", "123456", "Ленина"):
        assert sensitive not in serialised


def test_users_import_requires_email_column() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Фамилия", "Имя"])
    sheet.append(["Иванов", "Пётр"])
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    result = parse_catalog_file(output, "users", "users.xlsx")

    assert result.success is False
    assert any("Email" in error for error in result.errors)


def test_users_import_warns_on_duplicate_email() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Email", "Фамилия"])
    sheet.append(["dup@example.com", "Иванов"])
    sheet.append(["dup@example.com", "Петров"])
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    result = parse_catalog_file(output, "users", "users.xlsx")

    assert result.success is True
    assert len(result.data) == 2
    assert any("Дубль" in issue.problem for issue in result.issues)
